"""[live-e2e] Real ingest pipeline — Celery worker + MinerU sidecar + real model.

These hit LIVE external services (the Celery worker via Redis, the MinerU
sidecar, and the configured LLM) and cost real model tokens, so they only run
when ``WIKI_RUN_LIVE_E2E=1`` is set; otherwise they skip-with-reason. They lock
the full async dispatch path that the mock / in-process ``test_sources_ingest``
tests can't reach.

Validated 2026-06-12: worker completed the plan in ~11 s (1 edit, no error);
MinerU extracted text from a generated PDF into markdown.

Run them with:
    docker exec -e WIKI_RUN_LIVE_E2E=1 wiki-backend-1 python -m pytest tests/test_ingest_e2e.py -q
"""
from __future__ import annotations

import os
import time

import httpx
import pytest

LIVE = os.environ.get("WIKI_RUN_LIVE_E2E") == "1"
needs_live = pytest.mark.skipif(
    not LIVE,
    reason="live-e2e (worker + MinerU + real model, costs tokens) — set WIKI_RUN_LIVE_E2E=1",
)
BASE = os.environ.get("WIKI_TEST_BASE_URL", "http://localhost:8000")
ADMIN = {"X-User-Email": "live-e2e@example.com", "X-User-Role": "admin"}


def _make_pdf(text: str) -> bytes:
    """A minimal single-page text PDF (no 3rd-party PDF lib needed)."""
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
    ]
    stream = b"BT /F1 18 Tf 72 700 Td (" + text.encode() + b") Tj ET"
    objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = b"%PDF-1.4\n"
    offs = []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += str(i).encode() + b" 0 obj\n" + o + b"\nendobj\n"
    xp = len(out)
    out += b"xref\n0 " + str(len(objs) + 1).encode() + b"\n0000000000 65535 f \n"
    for off in offs:
        out += ("%010d 00000 n \n" % off).encode()
    out += (b"trailer\n<< /Size " + str(len(objs) + 1).encode()
            + b" /Root 1 0 R >>\nstartxref\n" + str(xp).encode() + b"\n%%EOF")
    return out


@needs_live
def test_ingest_worker_e2e_real_dispatch():
    """API -> Redis -> real Celery worker -> real model -> pending_review plan.

    The whole point is the *async dispatch*: we never touch the task body, we
    enqueue via the HTTP endpoint and poll until the live worker finishes."""
    src = (b"Hypatia of Alexandria (c. 350-415 CE) was a Neoplatonist philosopher, "
           b"astronomer, and mathematician, head of the Platonist school at Alexandria.")
    sid = None
    try:
        r = httpx.post(f"{BASE}/api/raw", headers=ADMIN,
                       files={"file": ("hypatia.txt", src, "text/plain")},
                       data={"title": "live-e2e Hypatia"}, timeout=30)
        assert r.status_code in (200, 201), r.text
        sid = r.json()["id"]

        t = httpx.post(f"{BASE}/api/raw/{sid}/ingest", headers=ADMIN, timeout=30)
        assert t.status_code in (200, 201), t.text
        rid = t.json()["id"]
        assert t.json()["status"] == "planning", "ingest should start in planning"

        run, status = None, "planning"
        deadline = time.time() + 180
        while time.time() < deadline:
            runs = httpx.get(f"{BASE}/api/raw/{sid}/runs", headers=ADMIN, timeout=15).json()
            run = next((x for x in runs if x["id"] == rid), None)
            status = run["status"] if run else "?"
            if status != "planning":
                break
            time.sleep(5)

        err = (run or {}).get("error") or ""
        if status == "failed" and ("rate_limit" in err.lower() or "429" in err):
            pytest.skip(f"model rate-limited (env tier too low): {err[:100]}")
        assert status == "pending_review", f"worker did not finish the plan: {status} err={err[:120]}"
        assert (run or {}).get("edits_count", 0) >= 1, "plan should propose at least one edit"
    finally:
        if sid is not None:
            httpx.request("DELETE", f"{BASE}/api/raw/{sid}", headers=ADMIN, timeout=20)


@needs_live
def test_mineru_pdf_conversion():
    """A generated text PDF is parsed by the real MinerU sidecar into markdown."""
    import asyncio

    from app.services.converters import convert_via_mineru

    pdf = _make_pdf("Tau scaling laws for Huawei Kirin chips and logic folding.")
    md = asyncio.run(convert_via_mineru(pdf, "tau.pdf", "application/pdf"))
    assert md, "MinerU returned no markdown (sidecar down or parse failed)"
    assert any(k in md.lower() for k in ("tau", "kirin", "scaling")), md[:200]


_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@needs_live
def test_mineru_office_conversion():
    """Real MinerU converts Office formats: xlsx -> a markdown table, docx ->
    text. (PDF is covered above; FR-ING-006 covers the dispatch with MinerU off.)
    Needs openpyxl + python-docx to author the fixtures."""
    import asyncio
    import io

    pytest.importorskip("openpyxl")
    pytest.importorskip("docx")
    from openpyxl import Workbook
    from docx import Document

    from app.services.converters import convert_via_mineru

    wb = Workbook()
    wb.active["A1"] = "Tau scaling law"
    wb.active["B1"] = "Kirin chips"
    xb = io.BytesIO(); wb.save(xb)
    md_x = asyncio.run(convert_via_mineru(xb.getvalue(), "x.xlsx", _XLSX))
    assert md_x and any(k in md_x.lower() for k in ("tau", "kirin")), (md_x or "")[:200]

    doc = Document(); doc.add_paragraph("Parmenides of Elea founded the Eleatic school.")
    db = io.BytesIO(); doc.save(db)
    md_d = asyncio.run(convert_via_mineru(db.getvalue(), "d.docx", _DOCX))
    assert md_d and "parmenides" in md_d.lower(), (md_d or "")[:200]
